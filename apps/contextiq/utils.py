import os
from PyPDF2 import PdfReader
import pdfplumber
from docx import Document as DocxDocument

def parse_pdf(file_path):
    """
    Parse PDF file (text + tables using pdfplumber) and return combined content as string.
    """
    content = ""

    # Extract text using PyPDF2
    try:
        reader = PdfReader(file_path)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                content += page_text + "\n"
    except Exception as e:
        print(f"Error reading PDF text: {e}")

    # Extract tables using pdfplumber
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                for table in page.extract_tables():
                    table_text = ""
                    for row in table:
                        row_text = [cell.strip() if cell else "" for cell in row]
                        table_text += "\t".join(row_text) + "\n"
                    content += "\n" + table_text + "\n"
    except Exception as e:
        print(f"Error parsing PDF tables: {e}")

    return content.strip()


def parse_docx(file_path):
    """
    Parse DOCX file (text + tables) and return combined content as string.
    """
    content = ""
    try:
        doc = DocxDocument(file_path)
        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                content += para.text + "\n"
        # Extract tables
        for table in doc.tables:
            table_text = ""
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                table_text += "\t".join(row_text) + "\n"
            content += "\n" + table_text + "\n"
    except Exception as e:
        print(f"Error parsing DOCX: {e}")

    return content.strip()


def parse_document(file_path):
    """
    Detect file type and parse accordingly.
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pdf':
        return parse_pdf(file_path)
    elif ext in ['.doc', '.docx']:
        return parse_docx(file_path)
    else:
        raise ValueError("Unsupported file type")
